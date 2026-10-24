# -*- coding: utf-8 -*-
"""
转换后校准（html-to-docx 的已知缺陷补丁）：
  1. run 级字体被统一写成 body 字体 -> 标题丢黑体；
  2. 西文/希腊字母/数学符号被套上中文字体（宋体、黑体自带西文字面，显得全角发扁）
     -> 统一改成 w:ascii / w:hAnsi / w:cs = Times New Roman，中文仍走 eastAsia；
  3. Heading 4 自带斜体 -> 关掉；
  4. 标题段前距偏小、表格跨页、题注孤立 -> 补齐。
"""
import os
import re
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

P = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "实验报告_Project2.docx")

TNR = "Times New Roman"
CJK_BODY = "宋体"
CJK_HEAD = "黑体"

SPACE_BEFORE = {"Heading 1": 12, "Heading 2": 12, "Heading 3": 10, "Heading 4": 8}
SPACE_AFTER = {"Heading 1": 6, "Heading 2": 6, "Heading 3": 4, "Heading 4": 3}


def set_run_fonts(rPr, cjk):
    """西文 = Times New Roman；中文 = 指定字体；去掉 hint 让 Word 按码位区自动分流。"""
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), TNR)
    rFonts.set(qn("w:hAnsi"), TNR)
    rFonts.set(qn("w:cs"), TNR)
    rFonts.set(qn("w:eastAsia"), cjk)
    h = qn("w:hint")
    if h in rFonts.attrib:
        del rFonts.attrib[h]


def clear_italic(rPr):
    for tag in ("w:i", "w:iCs"):
        e = rPr.find(qn(tag))
        if e is not None:
            rPr.remove(e)


def fix_paragraph(par, is_heading):
    for run in par.runs:
        rPr = run._element.get_or_add_rPr()
        set_run_fonts(rPr, CJK_HEAD if is_heading else CJK_BODY)
        if is_heading:
            clear_italic(rPr)
            run.font.italic = False


d = Document(P)
runs_fixed = 0

# --- 正文（含表格单元格）---
for par in d.paragraphs:
    is_h = par.style.name.startswith("Heading")
    if is_h:
        pf = par.paragraph_format
        pf.space_before = Pt(SPACE_BEFORE.get(par.style.name, 8))
        pf.space_after = Pt(SPACE_AFTER.get(par.style.name, 3))
        pf.line_spacing = 1.2
    fix_paragraph(par, is_h)
    runs_fixed += len(par.runs)

for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                fix_paragraph(par, par.style.name.startswith("Heading"))
                runs_fixed += len(par.runs)

# --- 页眉页脚 ---
for sec in d.sections:
    for part in (sec.header, sec.footer):
        try:
            for par in part.paragraphs:
                fix_paragraph(par, False)
        except Exception:
            pass

# --- 样式层 ---
for sname, cjk in [("Normal", CJK_BODY), ("Heading 1", CJK_HEAD), ("Heading 2", CJK_HEAD),
                   ("Heading 3", CJK_HEAD), ("Heading 4", CJK_HEAD)]:
    try:
        st = d.styles[sname]
        set_run_fonts(st.element.get_or_add_rPr(), cjk)
        if sname.startswith("Heading"):
            st.font.italic = False
            st.paragraph_format.space_before = Pt(SPACE_BEFORE[sname])
            st.paragraph_format.space_after = Pt(SPACE_AFTER[sname])
    except Exception as e:
        print("style err", sname, e)

# --- 小表尽量不跨页 ---
kept = 0
for t in d.tables:
    rows = t.rows
    if len(rows) > 10 or len(rows) < 2:
        continue
    for row in rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    kept += 1

# --- 题注/图片与后文同页 ---
glued = 0
for par in d.paragraphs:
    t = par.text.strip()
    has_img = par._element.findall(".//" + qn("w:drawing")) != []
    if has_img or re.match(r"^[图表]\s*\d+", t):
        par.paragraph_format.keep_with_next = True
        glued += 1

d.save(P)
print("runs fixed:", runs_fixed, "| tables kept:", kept, "| captions glued:", glued)
